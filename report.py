#!/usr/bin/env python3
"""
Campaign Excel Builder
======================
Usage:
    python campaign_excel_builder.py <file1.csv/xlsx> [file2 ...]

Steps:
    1. Run with your CSV/Excel file(s)
    2. Script prints unique campaign names
    3. You enter desired order (comma-separated numbers)
    4. Generates campaign_report_YYYY_MM.xlsx
"""

import sys, csv, re, json, argparse
from datetime import date, timedelta
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Config ────────────────────────────────────────────────────────────────────
COL_CAMPAIGN  = "Campaign name"
COL_DAY       = "Day"
COL_RESULTS   = "Results"
COL_COST      = "Cost per result"
COL_SPENT     = "Amount spent (INR)"

RED      = "C8102E"
WHITE    = "FFFFFF"
DARK_RED = "9B0E24"

# ── Styles ────────────────────────────────────────────────────────────────────
def red_cell(ws, row, col, value, bold=True, size=9):
    c = ws.cell(row=row, column=col, value=value)
    c.fill      = PatternFill("solid", fgColor=RED)
    c.font      = Font(color=WHITE, bold=bold, size=size, name="Arial")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return c

def plain_cell(ws, row, col, value):
    c = ws.cell(row=row, column=col, value=value)
    c.font      = Font(size=9, name="Arial")
    c.alignment = Alignment(horizontal="center", vertical="center")
    return c

def border_cell(c):
    thin = Side(style="thin", color="CCCCCC")
    c.border = Border(left=thin, right=thin, top=thin, bottom=thin)

# ── Parse files ───────────────────────────────────────────────────────────────
def parse_files(paths):
    rows = []
    for path in paths:
        if path.endswith(".csv"):
            with open(path, newline="", encoding="utf-8-sig") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    rows.append(row)
        else:
            import openpyxl as ox
            wb = ox.load_workbook(path, data_only=True)
            for ws in wb.worksheets:
                headers = None
                for r in ws.iter_rows(values_only=True):
                    if headers is None:
                        headers = [str(h).strip() if h else "" for h in r]
                    else:
                        rows.append(dict(zip(headers, r)))
    return rows

def unique_campaigns(rows):
    seen, out = set(), []
    for r in rows:
        name = str(r.get(COL_CAMPAIGN) or "").strip()
        if name and name != COL_CAMPAIGN and name not in seen:
            seen.add(name)
            out.append(name)
    return out

def parse_date(val):
    if not val: return None
    s = str(val).strip()
    # YYYY-MM-DD
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m: return date(int(m[1]), int(m[2]), int(m[3]))
    # DD/MM/YYYY
    m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m: return date(int(m[3]), int(m[2]), int(m[1]))
    return None

def build_lookup(rows, campaigns):
    camp_set = set(campaigns)
    # lookup[campaign][day_int] = {results, spent}
    lookup = {c: defaultdict(lambda: {"results": 0, "spent": 0.0}) for c in campaigns}

    for r in rows:
        # Normalize keys to lowercase for robust lookup
        r_norm = {str(k).lower().strip(): v for k, v in r.items()}
        
        name = str(r_norm.get(COL_CAMPAIGN.lower()) or "").strip()
        if name not in camp_set: continue
        
        d = parse_date(r_norm.get(COL_DAY.lower()))
        if not d: continue

        try: 
            results = float(r_norm.get(COL_RESULTS.lower()) or 0)
        except: 
            results = 0
            
        try: 
            # Try to get direct spend, fallback to results * cost
            spent = float(r_norm.get(COL_SPENT.lower()) or 0)
            if spent == 0:
                cost = float(r_norm.get(COL_COST.lower()) or 0)
                spent = results * cost
        except: 
            spent = 0.0

        lookup[name][d.day]["results"] += results
        lookup[name][d.day]["spent"]   += spent

    return lookup

def infer_month_year(rows):
    for r in rows:
        d = parse_date(r.get(COL_DAY))
        if d: return d.month, d.year
    t = date.today()
    return t.month, t.year

def days_in_month(month, year):
    if month == 12: return (date(year+1,1,1) - date(year,12,1)).days
    return (date(year, month+1, 1) - date(year, month, 1)).days

# ── Build Excel ───────────────────────────────────────────────────────────────
def build_excel(ordered, lookup, month, year, out_path):
    total_days = days_in_month(month, year)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Campaign Report"

    n = len(ordered)
    # Col layout: 1=Date, then per campaign: results col, cost col, then Total Spent
    # campaign i → results col = 2 + i*2, cost col = 3 + i*2
    # total spent col = 2 + n*2

    HEADER_ROW  = 1
    SUBHEAD_ROW = 2
    DATA_START  = 3
    DATA_END    = DATA_START + total_days - 1
    TOTAL_ROW   = DATA_END + 1

    total_col = 2 + n * 2

    def res_col(i): return 2 + i * 2
    def cpr_col(i): return 3 + i * 2

    # ── Row heights ───────────────────────────────────────────────────────────
    ws.row_dimensions[HEADER_ROW].height  = 40
    ws.row_dimensions[SUBHEAD_ROW].height = 20
    ws.row_dimensions[TOTAL_ROW].height   = 20
    for r in range(DATA_START, DATA_END + 1):
        ws.row_dimensions[r].height = 15

    # ── Col widths ────────────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 12
    for i in range(n):
        ws.column_dimensions[get_column_letter(res_col(i))].width = 10
        ws.column_dimensions[get_column_letter(cpr_col(i))].width = 14
    ws.column_dimensions[get_column_letter(total_col)].width = 13

    # ── Header row 1 ─────────────────────────────────────────────────────────
    # Date: merge A1:A2
    red_cell(ws, HEADER_ROW, 1, "Date")
    ws.merge_cells(start_row=HEADER_ROW, start_column=1,
                   end_row=SUBHEAD_ROW,  end_column=1)

    for i, name in enumerate(ordered):
        red_cell(ws, HEADER_ROW, res_col(i), name)
        ws.merge_cells(start_row=HEADER_ROW, start_column=res_col(i),
                       end_row=HEADER_ROW,   end_column=cpr_col(i))

    # Total Spent: merge row1:row2
    red_cell(ws, HEADER_ROW, total_col, "Total\nSpent")
    ws.merge_cells(start_row=HEADER_ROW, start_column=total_col,
                   end_row=SUBHEAD_ROW,  end_column=total_col)

    # ── Subheader row 2 ───────────────────────────────────────────────────────
    for i in range(n):
        red_cell(ws, SUBHEAD_ROW, res_col(i), "Results",         bold=False)
        red_cell(ws, SUBHEAD_ROW, cpr_col(i), "Cost per Result", bold=False)

    # ── Max day per campaign ──────────────────────────────────────────────────
    def max_day(name):
        days = list(lookup[name].keys())
        return max(days) if days else 0

    # ── Data rows ─────────────────────────────────────────────────────────────
    for d in range(1, total_days + 1):
        row = DATA_START + d - 1
        date_str = f"{d:02d}/{month:02d}/{year}"
        red_cell(ws, row, 1, date_str, bold=False, size=8)

        row_spent = 0.0
        has_data  = False

        for i, name in enumerate(ordered):
            mx  = max_day(name)
            day_data = lookup[name].get(d)

            if day_data:
                has_data = True
                results  = day_data["results"]
                spent    = day_data["spent"]
                cpr      = round(spent / results, 2) if results > 0 else 0

                plain_cell(ws, row, res_col(i), int(results) if results == int(results) else results)
                plain_cell(ws, row, cpr_col(i), cpr if cpr > 0 else "-")
                row_spent += spent
            elif d <= mx:
                plain_cell(ws, row, res_col(i), "-")
                plain_cell(ws, row, cpr_col(i), "-")
            # else: blank (beyond campaign's last active day)

        if has_data:
            plain_cell(ws, row, total_col, round(row_spent, 2))

        # light border for all data cells
        for col in range(1, total_col + 1):
            border_cell(ws.cell(row=row, column=col))

    # ── Total row ─────────────────────────────────────────────────────────────
    red_cell(ws, TOTAL_ROW, 1, "TOTAL")

    grand_total = 0.0
    for i, name in enumerate(ordered):
        total_results = sum(v["results"] for v in lookup[name].values())
        total_spent   = sum(v["spent"]   for v in lookup[name].values())
        grand_total  += total_spent
        
        # Show only Total Spent in the footer (per user request: no need for results 50, 43 etc)
        red_cell(ws, TOTAL_ROW, res_col(i), "")
        red_cell(ws, TOTAL_ROW, cpr_col(i), round(total_spent, 2))

    red_cell(ws, TOTAL_ROW, total_col, round(grand_total, 2))

    # ── Freeze panes ──────────────────────────────────────────────────────────
    ws.freeze_panes = ws.cell(row=DATA_START, column=2)

    wb.save(out_path)
    print(f"SUCCESS: {out_path}")

# ── Main ──────────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("files", nargs="*", help="File paths")
    parser.add_argument("--list", action="store_true", help="List unique campaigns as JSON")
    parser.add_argument("--order", help="Comma-separated indices or names for ordering")
    parser.add_argument("--out", help="Custom output filename")
    
    args = parser.parse_args()
    
    if not args.files:
        print("Usage: python report.py <files...> [--list] [--order 1,2,3]")
        sys.exit(0)

    rows = parse_files(args.files)
    campaigns = unique_campaigns(rows)
    
    if args.list:
        print(json.dumps(campaigns))
        return

    if args.order:
        try:
            # Try numeric indices first
            if "," in args.order and args.order.replace(",","").replace(" ","").isdigit():
                indices = [int(x.strip()) for x in args.order.split(",")]
                ordered = [campaigns[i-1] for i in indices if 0 < i <= len(campaigns)]
            else:
                # Treat as comma-separated names
                ordered = [x.strip() for x in args.order.split(",")]
        except:
            ordered = campaigns
    else:
        # Fallback to interactive if no --order or --list (for compatibility)
        if sys.stdin.isatty():
            print("\nUnique campaigns:")
            for i, c in enumerate(campaigns, 1):
                print(f"   {i}. {c}")
            ans = input(f"\nEnter order as comma-separated numbers (1-{len(campaigns)}), or Enter to keep as-is:\n   ").strip()
            if ans:
                indices = [int(x.strip()) for x in ans.split(",")]
                ordered = [campaigns[i-1] for i in indices]
            else:
                ordered = campaigns
        else:
            ordered = campaigns

    month, year = infer_month_year(rows)
    lookup = build_lookup(rows, ordered)
    out_path = args.out if args.out else f"campaign_report_{year}_{month:02d}.xlsx"
    
    build_excel(ordered, lookup, month, year, out_path)
    print(f"SUCCESS: {out_path}")

if __name__ == "__main__":
    main()